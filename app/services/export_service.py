import asyncio
import io
from collections.abc import Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from pydantic import BaseModel

from app.models.user import User
from app.services.report_service import ReportService

XLSX_MEDIA_TYPE = (
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
)

XLSX_FILENAMES = {
    "families": "families_export.xlsx",
    "members": "members_export.xlsx",
}


class ExportService:
    """Builds XLSX files from the role-scoped report pipeline.

    The report service already resolves MANAGER shelter scope and
    BLOCK_HEAD block scope; here we only convert the resulting rows into
    an .xlsx workbook. Workbook generation is CPU-bound, so it runs in a
    worker thread via ``asyncio.to_thread`` to keep the event loop free.
    """

    def __init__(self, report_service: ReportService):
        self.report_service = report_service

    async def export_families_xlsx(
        self,
        current_user: User,
        block_ids: list[int] | None = None,
        selected_ids: list[int] | None = None,
        columns: list[str] | None = None,
    ) -> bytes:
        rows = await self.report_service.get_families_report(
            current_user, block_ids, selected_ids
        )
        return await self._to_xlsx(rows, columns)

    async def export_members_xlsx(
        self,
        current_user: User,
        block_ids: list[int] | None = None,
        special_only: bool = False,
        selected_ids: list[int] | None = None,
        columns: list[str] | None = None,
    ) -> bytes:
        rows = await self.report_service.get_members_report(
            current_user, block_ids, special_only, selected_ids
        )
        return await self._to_xlsx(rows, columns)

    @staticmethod
    def _resolve_headers(
        rows: Sequence[BaseModel], columns: Sequence[str] | None
    ) -> list[str]:
        if not rows:
            return list(columns or [])
        all_keys = list(rows[0].model_dump().keys())
        if not columns:
            return all_keys
        return [col for col in columns if col in all_keys]  # Filter invalid columns

    async def _to_xlsx(
        self, rows: Sequence[BaseModel], columns: Sequence[str] | None
    ) -> bytes:
        headers = self._resolve_headers(rows, columns)
        rows_dump = [row.model_dump() for row in rows]
        return await asyncio.to_thread(self._build_workbook, headers, rows_dump)

    @staticmethod
    def _build_workbook(headers: list[str], rows: list[dict]) -> bytes:
        wb = Workbook()
        ws = wb.active
        ws.title = "export"

        if headers:
            ws.append(headers)
            header_fill = Font(bold=True)
            for cell in ws[1]:
                cell.font = header_fill

        for row in rows:
            ws.append([row.get(h) for h in headers])

        for idx, _header in enumerate(headers, start=1):
            column = get_column_letter(idx)
            widths = []
            for r in range(1, ws.max_row + 1):
                value = ws.cell(row=r, column=idx).value
                widths.append(len(str(value)) if value else 0)
            width = min(max(max(widths, default=0) + 2, 8), 40)
            ws.column_dimensions[column].width = width

        for row in ws.iter_rows():
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")

        ws.freeze_panes = "A2"

        buffer = io.BytesIO()
        wb.save(buffer)
        return buffer.getvalue()
"""Endpoint + service tests for the XLSX export pipeline.

- ExportService produces a valid .xlsx workbook (openpyxl-loadable)
- column selection filters headers and data
- members export contains one row per member
- GET /export/families returns the XLSX content type behind auth
"""

import io
from datetime import date

from fastapi import FastAPI
from httpx import ASGITransport, AsyncClient
from openpyxl import load_workbook
from sqlalchemy.ext.asyncio import AsyncSession

from app.api.deps import get_db, get_export_service
from app.core.security import create_access_token
from app.models.enums import (
    Gender,
    HousingType,
    MaritalStatus,
    ResidencyStatus,
    UserRole,
)
from app.repositories.auditRepository import AuditRepository
from app.repositories.familyRepository import FamilyRepository
from app.repositories.MemberRepository import MemberRepository
from app.schemas.family import FamilyCreate, MemberCreate
from app.services.audit_service import AuditService
from app.services.export_service import XLSX_MEDIA_TYPE, ExportService
from app.services.family_service import FamilyService
from app.services.report_service import ReportService

HEAD_ID = 400000006
SPOUSE_ID = 400000014


def _member(
    mid: int,
    name: str,
    gender: Gender = Gender.MALE,
    dob: date = date(1990, 1, 15),
) -> MemberCreate:
    return MemberCreate(
        id=mid,
        full_name=name,
        gender=gender,
        marital_status=MaritalStatus.MARRIED,
        date_of_birth=dob,
        relationship_to_head_id=1,
    )


def _family_create() -> FamilyCreate:
    return FamilyCreate(
        head_id=HEAD_ID,
        spouse_id=SPOUSE_ID,
        primary_phone_number="+970599123456",
        residency_status=ResidencyStatus.DISPLACED,
        housing_type=HousingType.TENT,
        original_city_id=1,
        current_shelter_center_id=1,
        shelter_block_id=1,
        shelter_quality_id=1,
        members=[
            _member(HEAD_ID, "Head of Family"),
            _member(SPOUSE_ID, "Spouse of Family", Gender.FEMALE, date(1992, 4, 22)),
        ],
    )


async def _seed_family(db: AsyncSession, admin_user) -> int:
    svc = FamilyService(
        FamilyRepository(db),
        MemberRepository(db),
        audit_service=AuditService(AuditRepository(db)),
    )
    family = await svc.create_family(_family_create(), admin_user)
    return family.id


def _export_service(db: AsyncSession) -> ExportService:
    return ExportService(
        ReportService(FamilyRepository(db), MemberRepository(db))
    )


def _make_app(db: AsyncSession) -> FastAPI:
    from app.api.v1.endpoints import export as export_endpoints

    async def _override_db():
        yield db

    app = FastAPI()
    app.include_router(export_endpoints.router, prefix="/export")
    app.dependency_overrides[get_db] = _override_db
    app.dependency_overrides[get_export_service] = lambda: _export_service(db)
    return app


# ── Service level ────────────────────────────────────────────────────────────


async def test_export_families_xlsx(db: AsyncSession, sample_lookups, admin_user):
    await _seed_family(db, admin_user)
    data = await _export_service(db).export_families_xlsx(admin_user)

    wb = load_workbook(io.BytesIO(data))
    ws = wb.active
    headers = [c.value for c in ws[1]]
    assert "family_id" in headers
    assert "member_count" in headers
    assert ws.max_row == 2  # header + 1 family row


async def test_export_families_column_selection(
    db: AsyncSession, sample_lookups, admin_user
):
    await _seed_family(db, admin_user)
    data = await _export_service(db).export_families_xlsx(
        admin_user, columns=["family_id", "head_name"]
    )

    wb = load_workbook(io.BytesIO(data))
    ws = wb.active
    assert [c.value for c in ws[1]] == ["family_id", "head_name"]
    assert ws["A2"].value is not None
    assert ws.max_row == 2


async def test_export_members_xlsx(db: AsyncSession, sample_lookups, admin_user):
    await _seed_family(db, admin_user)
    data = await _export_service(db).export_members_xlsx(admin_user)

    wb = load_workbook(io.BytesIO(data))
    ws = wb.active
    headers = [c.value for c in ws[1]]
    assert "member_id" in headers
    assert "member_name" in headers
    assert ws.max_row == 3  # header + 2 members


# ── Endpoint level ───────────────────────────────────────────────────────────


async def test_export_endpoint_returns_xlsx(
    db: AsyncSession, sample_lookups, admin_user
):
    await _seed_family(db, admin_user)
    token = create_access_token(
        {"sub": admin_user.username, "role": UserRole.SUPERADMIN.value}
    )
    app = _make_app(db)

    async with AsyncClient(
        transport=ASGITransport(app=app), base_url="http://test"
    ) as client:
        res = await client.get(
            "/export/families",
            headers={"Authorization": f"Bearer {token}"},
            params={"columns": ["family_id", "head_name"]},
        )

    assert res.status_code == 200
    assert res.headers["content-type"].startswith(XLSX_MEDIA_TYPE)
    wb = load_workbook(io.BytesIO(res.content))
    ws = wb.active
    assert [c.value for c in ws[1]] == ["family_id", "head_name"]


async def test_export_endpoint_requires_staff(
    db: AsyncSession, sample_lookups, admin_user
):
    await _seed_family(db, admin_user)
    app = _make_app(db)

    async with AsyncClient(
        transport=ASGITransport(app=app), base_url="http://test"
    ) as client:
        res = await client.get("/export/families")

    assert res.status_code == 401